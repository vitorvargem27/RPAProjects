from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
import os

#Pegando o arquivo do Excel
studentsExcelArchive = 'Alunos.xlsx'

#Pegando a planilha dos alunos
studentsSheet = load_workbook(studentsExcelArchive)

#Selecionando a aba da planilha
sheetSelected = studentsSheet['Nomes']

#fazendo um for da segunda linha até o quantidade de itens da Coluna A
for lineName in range(2, len(sheetSelected["A"]) + 1) :

    #abrindo o arquivo word de certificado padrão
    wordDocument = Document('certified1.docx')
    styleDocument = wordDocument.styles["Normal"]

    #o %s converte para string/texto de linha por linha e transformando em valor
    studentName = sheetSelected['A%s' % lineName].value

    for paragraph in wordDocument.paragraphs :
        if "@nome" in paragraph.text :
            paragraph.text = studentName
            font = styleDocument.font
            font.name = 'Agency FB'
            font.size = Pt(23)

    wordDocument.save(f'Certificado {studentName}.docx')
